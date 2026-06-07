
# -*- coding: utf-8 -*-
"""月度稳定性 — 用 openpyxl 追加，不破坏原格式"""
import pandas as pd, numpy as np, re, warnings, os, math
warnings.filterwarnings("ignore")

BINNING_PATH = r"C:\Users\6\Desktop\qiongdan数据2\分箱结果_头尾5.xlsx"
RAW_PATH = r"C:\Users\6\Desktop\qiongdan数据2\失联用户及黑名单变量样本.csv"
LABEL = "target3"
CUTOFF = "2026-05-01"

SUB_COLS = ['总人数', '逾期数', '坏率', '金额逾期率']

def parse_bin_boundary(bin_label):
    s = str(bin_label).strip()
    if any(k in s for k in ['特殊','空值','缺失','nan','NA','null']):
        return (None, None, True)
    parts = [p.strip() for p in s.split(',')]
    if len(parts) == 1:
        nums = re.findall(r'-?\d+\.?\d*', parts[0])
        return (float(nums[0]), float(nums[0]), False) if nums else (None, None, True)
    lo_str, hi_str = parts[0], parts[-1]
    lo = float('-inf') if 'inf' in lo_str.lower() else float(re.findall(r'-?\d+\.?\d*', lo_str)[0]) if re.findall(r'-?\d+\.?\d*', lo_str) else float('-inf')
    hi = float('inf') if 'inf' in hi_str.lower() else float(re.findall(r'-?\d+\.?\d*', hi_str)[0]) if re.findall(r'-?\d+\.?\d*', hi_str) else float('inf')
    return (lo, hi, s.startswith('['), s.endswith(']'), False)

def assign_bin(value, bin_defs):
    is_nan = value is None or (isinstance(value, float) and math.isnan(value))
    if is_nan:
        for bl, meta in bin_defs:
            if (len(meta)==3 and meta[2]==True) or (len(meta)==5 and meta[4]==True): return bl
        return None
    for bl, meta in bin_defs:
        if (len(meta)==3 and meta[2]==True) or (len(meta)==5 and meta[4]==True): continue
        if len(meta) == 3:
            lo, hi, spec = meta
            if lo == hi and abs(value-lo)<1e-9: return bl
            continue
        lo, hi, cl, cr, spec = meta
        if lo == hi and abs(value-lo)<1e-9: return bl
        left_ok = (value>=lo) if cl else (value>lo)
        if lo==float('-inf'): left_ok=True
        right_ok = (value<=hi) if cr else (value<hi)
        if hi==float('inf'): right_ok=True
        if left_ok and right_ok: return bl
    return None

print("读取数据...")
df = pd.read_csv(RAW_PATH, encoding='gbk')
df['overdue_flag'] = df[LABEL]
df = df[df['overdue_flag'].isin([0,1])].copy()
col_map = {c:c.replace('#','_') for c in df.columns if '#' in c}
df = df.rename(columns=col_map)
trn = df[df['apply_date'] < CUTOFF].copy()
tst = df[df['apply_date'] >= CUTOFF].copy()

scr = pd.read_excel(BINNING_PATH, sheet_name='训练筛选变量')
td = pd.read_excel(BINNING_PATH, sheet_name='Train分箱明细')
ted = pd.read_excel(BINNING_PATH, sheet_name='Test分箱明细')
td.rename(columns={'变量英文名':'feature','Bin':'bin_label'}, inplace=True, errors='ignore')
ted.rename(columns={'变量英文名':'feature','Bin':'bin_label'}, inplace=True, errors='ignore')
vars_list = scr['变量名称'].unique()
print(f"{len(vars_list)} 变量")

def build(data, detail, months):
    out = []
    for var in vars_list:
        vd = detail[detail['feature']==var]
        if len(vd)==0: continue
        if var not in data.columns: continue
        bdefs = [(r['bin_label'], parse_bin_boundary(r['bin_label'])) for _, r in vd.iterrows()]
        ds = data[[var,'apply_month','overdue_flag','money']].copy()
        ds['_bin'] = ds[var].apply(lambda v: assign_bin(v, bdefs))
        miss = ds[var].isna()
        if miss.any():
            sp = [bl for bl,m in bdefs if (len(m)==3 and m[2]==True) or (len(m)==5 and m[4]==True)]
            if sp: ds.loc[miss,'_bin'] = sp[0]
        grp = ds.groupby(['_bin','apply_month']).agg(obs=('overdue_flag','count'),bad=('overdue_flag','sum'),money=('money','sum')).reset_index()
        bdf = ds[ds['overdue_flag']==1].groupby(['_bin','apply_month'])['money'].sum().reset_index().rename(columns={'money':'bad_money'})
        grp = grp.merge(bdf, on=['_bin','apply_month'], how='left')
        grp['bad_money'] = grp['bad_money'].fillna(0.0)
        bo = vd['bin_label'].tolist()
        out.append(('VAR', var))
        out.append(('H1', months))
        out.append(('H2', None))
        cs = {m:{'obs':0,'bad':0,'money':0.0,'bad_money':0.0} for m in months}
        at = {'obs':0,'bad':0,'money':0.0,'bad_money':0.0}
        for bl in bo:
            rd = [str(bl)]
            rt = {'obs':0,'bad':0,'money':0.0,'bad_money':0.0}
            for m in months:
                sub = grp[(grp['_bin']==bl)&(grp['apply_month']==m)]
                obs = int(sub.iloc[0]['obs']) if len(sub) else 0
                bad = int(sub.iloc[0]['bad']) if len(sub) else 0
                money = float(sub.iloc[0]['money']) if len(sub) else 0.0
                bmoney = float(sub.iloc[0]['bad_money']) if len(sub) else 0.0
                rd += [obs, bad, bad/obs if obs else 0, bmoney/money if money else 0]
                rt['obs'] += obs; rt['bad'] += bad; rt['money'] += money; rt['bad_money'] += bmoney
                cs[m]['obs'] += obs; cs[m]['bad'] += bad; cs[m]['money'] += money; cs[m]['bad_money'] += bmoney
            tbr = rt['bad']/rt['obs'] if rt['obs'] else 0
            tmbr = rt['bad_money']/rt['money'] if rt['money'] else 0
            rd += [rt['obs'], rt['bad'], tbr, tmbr]
            for k in at: at[k] += rt[k]
            out.append(('DATA', rd))
        sr = ['合计']
        for m in months:
            br = cs[m]['bad']/cs[m]['obs'] if cs[m]['obs'] else 0
            mbr = cs[m]['bad_money']/cs[m]['money'] if cs[m]['money'] else 0
            sr += [cs[m]['obs'], cs[m]['bad'], br, mbr]
        tbr = at['bad']/at['obs'] if at['obs'] else 0
        tmbr = at['bad_money']/at['money'] if at['money'] else 0
        sr += [at['obs'], at['bad'], tbr, tmbr]
        out.append(('SUM', sr))
        out.append(('BLK', None))
    return out

tm = sorted(trn['apply_month'].unique())
tem = sorted(tst['apply_month'].unique())
tr = build(trn, td, tm)
ter = build(tst, ted, tem)

print("写入 Excel（openpyxl 追加，保留原格式）...")
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import DataBarRule, CellIsRule

wb = openpyxl.load_workbook(BINNING_PATH)

def write_stability_sheet(ws, rows, months):
    nm = len(months)
    wt = 1 + nm*4 + 4

    hf = PatternFill(start_color='366092',end_color='366092',fill_type='solid')
    hfn = Font(bold=True,color='FFFFFF',size=10)
    hf2 = PatternFill(start_color='5B9BD5',end_color='5B9BD5',fill_type='solid')
    hfn2 = Font(bold=True,color='FFFFFF',size=9)
    vf = Font(bold=True,size=11,color='366092')
    tb = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
    nf = '#,##0'
    pf = '0.00%'
    sf = PatternFill(start_color='D9E2F3',end_color='D9E2F3',fill_type='solid')
    sfn = Font(bold=True)

    ws.column_dimensions['A'].width = 18
    for ci in range(1, wt):
        ws.column_dimensions[get_column_letter(ci+1)].width = 10

    ri = 1  # openpyxl is 1-indexed
    data_start = -1
    bar_ranges = []

    for item in rows:
        t, d = item
        if t == 'VAR':
            ws.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=wt)
            c = ws.cell(row=ri, column=1, value=f"变量: {d}")
            c.font = vf
            c.border = Border(bottom=Side(style='thin'))
            ri += 1
        elif t == 'H1':
            c = ws.cell(row=ri, column=1, value='箱')
            c.fill = hf; c.font = hfn; c.alignment = Alignment(horizontal='center'); c.border = tb
            for mi, m in enumerate(months):
                cs = 2 + mi*4
                ws.merge_cells(start_row=ri, start_column=cs, end_row=ri, end_column=cs+3)
                c = ws.cell(row=ri, column=cs, value=m)
                c.fill = hf; c.font = hfn; c.alignment = Alignment(horizontal='center'); c.border = tb
            cs = 2 + nm*4
            ws.merge_cells(start_row=ri, start_column=cs, end_row=ri, end_column=cs+3)
            c = ws.cell(row=ri, column=cs, value='合计')
            c.fill = hf; c.font = hfn; c.alignment = Alignment(horizontal='center'); c.border = tb
            ri += 1
            # Sub-headers
            c = ws.cell(row=ri, column=1)
            c.fill = hf2; c.border = tb
            for mi in range(nm):
                for si, s in enumerate(SUB_COLS):
                    c = ws.cell(row=ri, column=2+mi*4+si, value=s)
                    c.fill = hf2; c.font = hfn2; c.alignment = Alignment(horizontal='center'); c.border = tb
            for si, s in enumerate(SUB_COLS):
                c = ws.cell(row=ri, column=2+nm*4+si, value=s)
                c.fill = hf2; c.font = hfn2; c.alignment = Alignment(horizontal='center'); c.border = tb
            ri += 1
            data_start = ri
        elif t == 'DATA':
            for ci, v in enumerate(d):
                c = ws.cell(row=ri, column=ci+1)
                if ci == 0:
                    c.value = v; c.number_format = nf
                else:
                    ct = (ci-1) % 4
                    if ct in [0, 1]:
                        c.value = v; c.number_format = nf
                    else:
                        c.value = v; c.number_format = pf
                c.border = tb; c.alignment = Alignment(horizontal='center')
            ri += 1
        elif t == 'SUM':
            data_end = ri - 1
            if data_end >= data_start:
                for mi in range(nm + 1):
                    for si in [2, 3]:  # 坏率=2, 金额逾期率=3
                        col = 2 + mi*4 + si
                        bar_ranges.append((col, data_start, data_end))
            for ci, v in enumerate(d):
                c = ws.cell(row=ri, column=ci+1)
                if ci == 0:
                    c.value = v; c.font = sfn; c.alignment = Alignment(horizontal='center')
                else:
                    ct = (ci-1) % 4
                    if ct in [0, 1]:
                        c.value = v; c.number_format = nf
                    else:
                        c.value = v; c.number_format = pf
                c.fill = sf; c.font = sfn; c.border = tb; c.alignment = Alignment(horizontal='center')
            ri += 1
            data_start = -1
        elif t == 'BLK':
            ri += 1

    # Data bars
    for col, sr, er in bar_ranges:
        cl = get_column_letter(col)
        rule = DataBarRule(start_type='min', end_type='max', color='5B9BD5', showValue=True)
        ws.conditional_formatting.add(f'{cl}{sr}:{cl}{er}', rule)

    ws.freeze_panes = 'B2'

# Remove old stability sheets if exist, create new
for s_name in ['Train月度稳定性', 'Test月度稳定性']:
    if s_name in wb.sheetnames:
        del wb[s_name]

ws_tr = wb.create_sheet('Train月度稳定性')
write_stability_sheet(ws_tr, tr, tm)
print(f"  Train月度稳定性: {len(tr)} rows (含空行)")

ws_te = wb.create_sheet('Test月度稳定性')
write_stability_sheet(ws_te, ter, tem)
print(f"  Test月度稳定性:  {len(ter)} rows (含空行)")

wb.save(BINNING_PATH)
wb.close()
print(f"完成! -> {BINNING_PATH}")
print(f"  Train月度稳定性: {len(vars_list)} 变量, 月份: {tm}")
print(f"  Test月度稳定性:  {len(vars_list)} 变量, 月份: {tem}")
